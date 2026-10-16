"""
laporan_admin.py — Halaman Admin "Data Laporan": seluruh hasil proses & abnormal.

Fitur:
  * Filter tanggal, status, kondisi, operator, part
  * Ringkasan angka (jumlah laporan, ACT, NG, rasio NG, rata-rata %Prod, jam kerja)
  * Tab: Hasil Proses · Abnormal · Rekap Operator · Rekap Part
  * Download 1 file Excel (.xlsx) berisi semua tabel di sheet terpisah
Hanya membaca data; tidak ada yang diubah atau dihapus.
"""
from __future__ import annotations

from datetime import timedelta

import pandas as pd
import streamlit as st

import db
import ui
from db import DBError, now_wib

TZ = "Asia/Jakarta"


def _jam(series: pd.Series) -> pd.Series:
    t = pd.to_datetime(series, utc=True, errors="coerce").dt.tz_convert(TZ)
    return t.dt.strftime("%H:%M").fillna("")


def _angka(df: pd.DataFrame, kolom: list[str]) -> pd.DataFrame:
    for k in kolom:
        if k in df.columns:
            df[k] = pd.to_numeric(df[k], errors="coerce")
    return df


def _cocok(df: pd.DataFrame, kolom: list[str], kata: str) -> pd.Series:
    kata = kata.strip().upper()
    if not kata:
        return pd.Series(True, index=df.index)
    hasil = pd.Series(False, index=df.index)
    for k in kolom:
        if k in df.columns:
            hasil |= df[k].astype(str).str.upper().str.contains(kata, regex=False, na=False)
    return hasil


def tampilkan() -> None:
    hari_ini = now_wib().date()

    # ---------------- Header ----------------
    c_judul, c_aksi = st.columns([3, 2], vertical_alignment="center")
    with c_judul:
        st.markdown('<div class="sec">Admin</div><div class="bigname">📊 Data Laporan Produksi</div>',
                    unsafe_allow_html=True)
    with c_aksi:
        b1, b2 = st.columns(2)
        if b1.button("↻ Muat Ulang", width="stretch"):
            db.get_laporan_proses.clear()
            db.get_laporan_abnormal.clear()
            st.rerun()
        if b2.button("← Kembali", width="stretch"):
            st.session_state.pop("halaman", None)
            st.rerun()

    # ---------------- Filter ----------------
    with st.container(key="card_filter"):
        f1, f2 = st.columns([2, 3], gap="medium")
        with f1:
            rentang = st.date_input("Rentang tanggal", value=(hari_ini - timedelta(days=6), hari_ini),
                                    max_value=hari_ini, format="DD/MM/YYYY", key="adm_rentang")
            status = st.radio("Status", ["Semua", "FINISH", "START"], horizontal=True, key="adm_status")
        with f2:
            g1, g2 = st.columns(2)
            cari_op = g1.text_input("Operator / NIK", key="adm_op")
            cari_part = g2.text_input("Part No / Nama Part", key="adm_part")
            kondisi = st.radio("Kondisi", ["Semua", "Normal", "Overtime"], horizontal=True, key="adm_kondisi")

    if not isinstance(rentang, (list, tuple)) or len(rentang) != 2:
        st.info("Pilih tanggal awal dan akhir.")
        return
    tgl_awal, tgl_akhir = rentang

    try:
        with st.spinner("Mengambil data..."):
            proses = pd.DataFrame(db.get_laporan_proses(tgl_awal, tgl_akhir))
            abnormal = pd.DataFrame(db.get_laporan_abnormal(tgl_awal, tgl_akhir))
    except DBError as e:
        st.error(f"❌ {e}")
        return

    # ---------------- Olah data proses ----------------
    if not proses.empty:
        proses = _angka(proses, ["act", "ng", "persen_prod", "rasio_ng", "total_jam", "total_istirahat", "sec_pcs"])
        mask = _cocok(proses, ["nama", "nik"], cari_op) & _cocok(proses, ["part_no", "part_name"], cari_part)
        if status != "Semua":
            mask &= proses["status"] == status
        if kondisi != "Semua":
            mask &= proses["condition"] == kondisi
        proses = proses[mask]

    if not abnormal.empty:
        abnormal = _angka(abnormal, ["total_waktu"])
        abnormal = abnormal[_cocok(abnormal, ["operator", "nik"], cari_op)
                            & _cocok(abnormal, ["part_no", "part_name"], cari_part)]

    # ---------------- Susun semua tabel (dipakai tampilan & Excel) ----------------
    selesai = proses[proses["status"] == "FINISH"] if not proses.empty else proses
    total_act = int(selesai["act"].sum()) if not selesai.empty else 0
    total_ng = int(selesai["ng"].sum()) if not selesai.empty else 0
    rata_prod = selesai["persen_prod"].dropna().mean() if not selesai.empty else float("nan")
    berjalan = int((proses["status"] == "START").sum()) if not proses.empty else 0
    total_jam = float(selesai["total_jam"].sum()) if not selesai.empty else 0.0
    menit_abn = int(abnormal["total_waktu"].sum()) if not abnormal.empty else 0

    tabel = _tabel_proses(proses)
    tabel_ab = _tabel_abnormal(abnormal)
    rekap_kode = _rekap_kode(abnormal)
    rekap_op = _rekap(selesai, ["nama", "nik"], {"nama": "Operator", "nik": "NIK"})
    rekap_part = _rekap(selesai, ["part_no", "part_name", "urutan_proses"],
                        {"part_no": "Part No", "part_name": "Part Name", "urutan_proses": "Urutan"})

    # ---------------- Ringkasan ----------------
    ui.kpi_grid([
        ("Laporan Selesai", f"{len(selesai)}", "", f"{tgl_awal:%d/%m} – {tgl_akhir:%d/%m/%Y}"),
        ("Sedang Berjalan", f"{berjalan}", "accent" if berjalan else ""),
        ("Total ACT", f"{total_act:,}".replace(",", "."), "ok"),
        ("Total NG", f"{total_ng:,}".replace(",", "."), "bad" if total_ng else ""),
        ("Rasio NG", f"{total_ng / total_act * 100:.2f}%" if total_act else "—", ""),
        ("Rata-rata %Prod", f"{rata_prod:.1f}%" if pd.notna(rata_prod) else "—", "accent"),
        ("Total Jam Kerja", f"{total_jam:.1f} j" if not selesai.empty else "—", ""),
        ("Total Abnormal", f"{menit_abn} m", "bad" if not abnormal.empty else "", f"{len(abnormal)} catatan"),
    ])

    # ---------------- Download: 1 file Excel, banyak sheet ----------------
    ringkasan = [
        ("Periode", f"{tgl_awal:%d/%m/%Y} – {tgl_akhir:%d/%m/%Y}"),
        ("Filter Status", status),
        ("Filter Kondisi", kondisi),
        ("Filter Operator / NIK", cari_op or "-"),
        ("Filter Part", cari_part or "-"),
        ("Dibuat", now_wib().strftime("%d/%m/%Y %H:%M WIB")),
        ("", ""),
        ("Laporan Selesai", len(selesai)),
        ("Sedang Berjalan", berjalan),
        ("Total ACT", total_act),
        ("Total NG", total_ng),
        ("Rasio NG (%)", round(total_ng / total_act * 100, 2) if total_act else None),
        ("Rata-rata % Prod", round(float(rata_prod), 2) if pd.notna(rata_prod) else None),
        ("Total Jam Kerja", round(total_jam, 2)),
        ("Catatan Abnormal", len(abnormal)),
        ("Total Menit Abnormal", menit_abn),
    ]
    try:
        isi_excel = buat_excel(ringkasan, {
            "Hasil Proses":    tabel,
            "Abnormal":        tabel_ab,
            "Rekap Abnormal":  rekap_kode,
            "Rekap Operator":  rekap_op,
            "Rekap Part":      rekap_part,
        })
        st.download_button(
            "⬇️  DOWNLOAD EXCEL (SEMUA DATA)", data=isi_excel, key="go_excel",
            file_name=f"Laporan_Produksi_Press_{tgl_awal:%Y%m%d}_{tgl_akhir:%Y%m%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
        )
        st.caption("Satu file berisi 6 sheet: Ringkasan · Hasil Proses · Abnormal · Rekap Abnormal · "
                   "Rekap Operator · Rekap Part — sesuai filter di atas.")
    except Exception as e:
        st.error(f"Gagal membuat file Excel: {e}")

    # ---------------- Tampilan per tab (tetap terpisah) ----------------
    tab_p, tab_a, tab_o, tab_part = st.tabs(["📋 Hasil Proses", "⚠️ Abnormal", "👷 Rekap Operator", "🔩 Rekap Part"])
    fmt_persen = {
        "% Prod":           st.column_config.NumberColumn(format="%.2f%%"),
        "Rasio NG %":       st.column_config.NumberColumn(format="%.2f%%"),
        "Rata-rata % Prod": st.column_config.NumberColumn(format="%.1f%%"),
        "Total Jam":        st.column_config.NumberColumn(format="%.2f"),
        "Sec/Pcs":          st.column_config.NumberColumn(format="%.2f"),
    }

    with tab_p:
        if tabel.empty:
            st.info("Belum ada data proses untuk filter ini.")
        else:
            st.dataframe(tabel, hide_index=True, width="stretch", height=520, column_config=fmt_persen)

    with tab_a:
        if tabel_ab.empty:
            st.info("Belum ada data abnormal untuk filter ini.")
        else:
            c_grafik, c_rekap = st.columns([3, 2], gap="medium")
            with c_grafik:
                st.bar_chart(rekap_kode, x="Kode", y="Total Menit", color=ui.ACCENT, height=280)
            with c_rekap:
                st.dataframe(rekap_kode, hide_index=True, width="stretch", height=280)
            st.dataframe(tabel_ab, hide_index=True, width="stretch", height=420)

    with tab_o:
        if rekap_op.empty:
            st.info("Belum ada laporan FINISH untuk filter ini.")
        else:
            st.dataframe(rekap_op, hide_index=True, width="stretch", column_config=fmt_persen)

    with tab_part:
        if rekap_part.empty:
            st.info("Belum ada laporan FINISH untuk filter ini.")
        else:
            st.dataframe(rekap_part, hide_index=True, width="stretch", column_config=fmt_persen)


# =====================================================================
# PENYUSUN TABEL
# =====================================================================
KOLOM_PROSES = ["Tanggal", "Start", "Finish", "Operator", "NIK", "Part No", "Part Name", "Model", "Urutan",
                "Line", "Kondisi", "Sec/Pcs", "ACT", "NG", "% Prod", "Rasio NG %", "Istirahat (m)",
                "Total Jam", "Status"]
KOLOM_ABNORMAL = ["Tanggal", "Jam", "Operator", "Mesin", "Part No", "Part Name", "Urutan", "Kode",
                  "Uraian", "Menit", "Keterangan"]
KOLOM_REKAP = ["Laporan", "ACT", "NG", "Rasio NG %", "Rata-rata % Prod", "Total Jam"]


def _tabel_proses(proses: pd.DataFrame) -> pd.DataFrame:
    if proses.empty:
        return pd.DataFrame(columns=KOLOM_PROSES)
    return pd.DataFrame({
        "Tanggal":       proses["tanggal"],
        "Start":         _jam(proses["actual_time_start"]),
        "Finish":        _jam(proses["actual_time_finish"]),
        "Operator":      proses["nama"],
        "NIK":           proses["nik"],
        "Part No":       proses["part_no"],
        "Part Name":     proses["part_name"],
        "Model":         proses["model"],
        "Urutan":        proses["urutan_proses"],
        "Line":          proses["actual_line"],
        "Kondisi":       proses["condition"],
        "Sec/Pcs":       proses["sec_pcs"],
        "ACT":           proses["act"],
        "NG":            proses["ng"],
        "% Prod":        proses["persen_prod"],
        "Rasio NG %":    proses["rasio_ng"],
        "Istirahat (m)": proses["total_istirahat"],
        "Total Jam":     proses["total_jam"],
        "Status":        proses["status"],
    })


def _tabel_abnormal(abnormal: pd.DataFrame) -> pd.DataFrame:
    if abnormal.empty:
        return pd.DataFrame(columns=KOLOM_ABNORMAL)
    return pd.DataFrame({
        "Tanggal":    abnormal["tanggal"],
        "Jam":        abnormal["jam_input"].astype(str).str[:5],
        "Operator":   abnormal["operator"],
        "Mesin":      abnormal["mesin"],
        "Part No":    abnormal["part_no"],
        "Part Name":  abnormal["part_name"],
        "Urutan":     abnormal["urutan_proses"],
        "Kode":       abnormal["kode_abnormal"],
        "Uraian":     abnormal["uraian_abnormal"],
        "Menit":      abnormal["total_waktu"],
        "Keterangan": abnormal["keterangan"],
    })


def _rekap_kode(abnormal: pd.DataFrame) -> pd.DataFrame:
    if abnormal.empty:
        return pd.DataFrame(columns=["Kode", "Uraian", "Kejadian", "Total Menit"])
    return (abnormal.groupby(["kode_abnormal", "uraian_abnormal"], dropna=False)["total_waktu"]
            .agg(["count", "sum"]).reset_index()
            .rename(columns={"kode_abnormal": "Kode", "uraian_abnormal": "Uraian",
                             "count": "Kejadian", "sum": "Total Menit"})
            .sort_values("Total Menit", ascending=False))


def _rekap(selesai: pd.DataFrame, kunci: list[str], nama_kolom: dict) -> pd.DataFrame:
    if selesai.empty:
        return pd.DataFrame(columns=list(nama_kolom.values()) + KOLOM_REKAP)
    r = (selesai.groupby(kunci, dropna=False)
         .agg(Laporan=("id", "count"), ACT=("act", "sum"), NG=("ng", "sum"),
              Rata_Prod=("persen_prod", "mean"), Total_Jam=("total_jam", "sum"))
         .reset_index())
    r["Rasio NG %"] = (r["NG"] / r["ACT"].where(r["ACT"] > 0) * 100).round(2)
    r["Rata_Prod"] = r["Rata_Prod"].round(2)
    r = r.rename(columns={**nama_kolom, "Rata_Prod": "Rata-rata % Prod", "Total_Jam": "Total Jam"})
    return r[list(nama_kolom.values()) + KOLOM_REKAP].sort_values("ACT", ascending=False)


# =====================================================================
# EXCEL
# =====================================================================
FORMAT_ANGKA = {
    "% Prod": '0.00"%"', "Rasio NG %": '0.00"%"', "Rata-rata % Prod": '0.00"%"',
    "Total Jam": "0.00", "Sec/Pcs": "0.00", "Tanggal": "dd/mm/yyyy",
    "ACT": "#,##0", "NG": "#,##0", "Laporan": "#,##0", "Menit": "#,##0",
    "Total Menit": "#,##0", "Kejadian": "#,##0", "Istirahat (m)": "#,##0",
}


def buat_excel(ringkasan: list[tuple], sheets: dict[str, pd.DataFrame]) -> bytes:
    """Satu workbook .xlsx: sheet Ringkasan + satu sheet per tabel, dengan header, filter & lebar kolom."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", start_color="1F2A37")
    head_font = Font(bold=True, color="FFFFFF")
    judul_font = Font(bold=True, size=14)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Laporan Produksi Dept. Press — PT Indosafety Sentosa"
    ws["A1"].font = judul_font
    for i, (label, nilai) in enumerate(ringkasan, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=bool(label))
        ws.cell(row=i, column=2, value=nilai).alignment = Alignment(horizontal="left")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34

    for nama, df in sheets.items():
        ws = wb.create_sheet(nama[:31])
        df = df.copy()
        if "Tanggal" in df.columns:
            df["Tanggal"] = pd.to_datetime(df["Tanggal"], errors="coerce").dt.date
        kolom = list(df.columns)
        ws.append(kolom)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for baris in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NA else v for v in baris])

        for j, k in enumerate(kolom, start=1):
            huruf = get_column_letter(j)
            fmt = FORMAT_ANGKA.get(k)
            if fmt:
                for (sel,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    sel.number_format = fmt
            panjang = max([len(str(k))] + [len(str(v)) for v in df[k].head(500) if v is not None]) if len(df) else len(k)
            ws.column_dimensions[huruf].width = min(max(10, panjang + 2), 45)
        ws.freeze_panes = "A2"
        if kolom:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{max(ws.max_row, 1)}"
        if df.empty:
            ws.cell(row=2, column=1, value="(tidak ada data untuk filter ini)").font = Font(italic=True, color="888888")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
